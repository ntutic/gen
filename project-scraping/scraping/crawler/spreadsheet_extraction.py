"""Read a deliberately selected XLSX worksheet from a Scrapy response."""

from io import BytesIO


def xlsx_rows(response, *, sheet_name: str) -> list[tuple]:
    """Preserve cell types and blanks; spiders validate headers, identities and totals.

    Formulas are never executed. Only the publisher's cached formula values are
    read; a missing cache remains None and must not become a fabricated zero.
    """
    from openpyxl import load_workbook

    if not sheet_name or not response.body.startswith(b"PK"):
        raise ValueError("Expected XLSX response bytes and an explicit worksheet name")
    workbook = load_workbook(BytesIO(response.body), read_only=True, data_only=True, keep_links=False)
    try:
        if sheet_name not in workbook.sheetnames:
            raise ValueError(f"XLSX worksheet {sheet_name!r} is missing")
        rows = list(workbook[sheet_name].iter_rows(values_only=True))
        if not rows or not any(any(value is not None for value in row) for row in rows):
            raise ValueError(f"XLSX worksheet {sheet_name!r} is empty")
        return rows
    finally:
        workbook.close()
